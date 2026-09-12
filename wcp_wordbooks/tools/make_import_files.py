# -*- coding: utf-8 -*-
"""生成 WCP 游戏官方导入通道所需的配套文件:
  - 每级别 .xlsx (游戏内可用 Excel 导入创建词书)
  - 全级别 SQLite 外接词库 wcp_jlpt.db (pron 表: word/meaning)

游戏内 Excel 导入契约 (逆向自 SaveSourceType, NPOI):
  - 读第一个 sheet, 从第 0 行开始逐行读 A/B 两列
  - A列=单词, B列=释义(游戏内显示文本)  -> 因此【不要加表头行】
  - C列以后被忽略, 可放读音/例句等人工参考信息
外接词库通道: persistentDataPath 下 .db, 表名/列名在游戏 UI 里指定,
  默认兼容 pron(word, meaning)。
"""
import json
import sqlite3
import hashlib
import unicodedata
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'output'
IMPORT = OUT / 'import'

# A=单词 B=释义(游戏读取); C-F 仅人工参考
SLOT_MAP = {
    'JLPT_N5N4_初级': ['n5', 'n4'],
    'JLPT_N3': ['n3'],
    'JLPT_N2': ['n2'],
    'JLPT_N1': ['n1'],
}
COMBINED_NAME = '日语词库(猫条版)'


def fingerprint(words):
    """Match BookProfiles.cs: normalized, sorted full word set SHA-256."""
    normalized = sorted(unicodedata.normalize('NFC', w.strip()) for w in words)
    payload = ''.join(w + '\n' for w in normalized).encode('utf-8')
    return hashlib.sha256(payload).hexdigest()


def save_book(path, rows):
    """Write WCP's no-header import workbook; only columns A/B are game data."""
    wb = Workbook()
    ws = wb.active
    ws.title = '词汇表'
    for row in rows:
        ws.append(row)
    for col, w in zip('ABCDEF', (18, 46, 14, 40, 44, 8)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(vertical='center', wrap_text=True)
    wb.save(path)


def main():
    combined_only = '--combined-only' in sys.argv
    data = json.loads((OUT / 'jlpt_books.json').read_text(encoding='utf-8'))
    levels = data['levels']
    IMPORT.mkdir(parents=True, exist_ok=True)

    all_rows = []
    for name, lvs in SLOT_MAP.items():
        count = 0
        rows = []
        for lv in lvs:
            for r in levels[lv]:
                meaning = r.get('meaning') or r.get('meaning_en') or ''
                if not r['word'] or not meaning:
                    continue
                # 游戏: A=单词, B=释义; C/D/E=读音/英文/例句(仅人工参考)
                row = [r['word'], meaning, r['reading'], r['meaning_en'],
                       r['example_ja'], r['level']]
                rows.append(row)
                all_rows.append(row)
                count += 1
        if not combined_only:
            path = IMPORT / f'{name}.xlsx'
            save_book(path, rows)
            print(f'{path.name}: {count} 词')

    # 可移植单册：四个级别按 N5 → N1 合并，同词只保留第一次出现的释义。
    # 游戏导入不需要表头；完整词形签名写入 manifest，供 BepInEx 插件精确识别。
    seen_combined = set()
    combined_rows = []
    for row in all_rows:
        word = row[0]
        if word in seen_combined:
            continue
        seen_combined.add(word)
        combined_rows.append(row)
    combined_path = IMPORT / f'{COMBINED_NAME}.xlsx'
    save_book(combined_path, combined_rows)
    combined_manifest = {
        'id': 'catbar-jlpt-complete',
        'language': 'ja',
        'display_name': COMBINED_NAME,
        'word_count': len(combined_rows),
        'fingerprint_sha256': fingerprint([r[0] for r in combined_rows]),
        'source': 'JLPT N5, N4, N3, N2, N1; duplicate word forms removed',
    }
    (IMPORT / '日语词库(猫条版).profile.json').write_text(
        json.dumps(combined_manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'{combined_path.name}: {len(combined_rows)} 去重词; sha256={combined_manifest["fingerprint_sha256"]}')

    if combined_only:
        return

    # 外接词库 SQLite
    db_path = IMPORT / 'wcp_jlpt.db'
    if db_path.exists():
        db_path.unlink()
    con = sqlite3.connect(db_path)
    cur = con.cursor()
    cur.execute('CREATE TABLE pron (word TEXT PRIMARY KEY, meaning TEXT)')
    seen = set()
    for word, meaning, *_rest in all_rows:
        if word in seen:
            continue
        seen.add(word)
        cur.execute('INSERT INTO pron VALUES (?, ?)', (word, meaning))
    cur.execute('''CREATE TABLE jlpt_all (
        word TEXT, meaning TEXT, reading TEXT, meaning_en TEXT,
        example_ja TEXT, level TEXT)''')
    for row in all_rows:
        cur.execute('INSERT INTO jlpt_all VALUES (?,?,?,?,?,?)', row)
    con.commit()
    con.close()
    print(f'{db_path.name}: pron {len(seen)} 词 (去重), jlpt_all {len(all_rows)} 行')


if __name__ == '__main__':
    main()
