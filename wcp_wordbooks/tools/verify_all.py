# -*- coding: utf-8 -*-
"""全面验证: 例句生产/词元数据/翻译/游戏DB/音频/导入文件 是否有遗漏。只读。"""
import json
import re
import sqlite3
import sys
from collections import Counter
from pathlib import Path

import wcp_paths

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parent.parent
WORK = ROOT / 'work'
OUT = ROOT / 'output'
SA = wcp_paths.streaming_assets()
FULL = wcp_paths.full_db()
ONLY = wcp_paths.only_db()
VOC = (Path.home() / 'AppData' / 'LocalLow' / 'WCP' /
       'packs' / 'ja' / 'audio' / 'word')

print('=' * 60)
print('[1] 例句产出: 词覆盖与条数')
need = {}   # word -> reading/meaning/level
for n in range(40):
    cf = WORK / f'gen_words_{n:02d}.json'
    if not cf.exists():
        continue
    for w, v in json.loads(cf.read_text(encoding='utf-8')).items():
        need[w] = v
gen = {}
for p in sorted(WORK.glob('gen_out_*.json')):
    try:
        gen.update(json.loads(p.read_text(encoding='utf-8')))
    except Exception as e:
        print('  !! 解析失败:', p.name, e)
cnt = Counter(len(gen.get(w, [])) for w in need)
print(f'  词表词数 {len(need)}, 产出覆盖 {len(gen)}')
print(f'  每词例句数分布: {dict(sorted(cnt.items()))}')
bad = [w for w in need if len(gen.get(w, [])) != 3]
print(f'  不等于3条的词: {len(bad)}', bad[:10])

print('[2] 词元数据: reading/meaning 缺失')
no_rd = [w for w, v in need.items() if not (v.get('reading') or '').strip()]
no_mn = [w for w, v in need.items() if not (v.get('meaning') or '').strip()]
kana_word = re.compile(r'^[\u3040-\u30ffー・A-Za-z0-9Ａ-Ｚａ-ｚ０-９]+$')
no_rd_kanji = [w for w in no_rd if not kana_word.match(w)]
print(f'  reading 缺失 {len(no_rd)} (其中含汉字词 {len(no_rd_kanji)})',
      no_rd_kanji[:10])
print(f'  meaning 缺失 {len(no_mn)}', no_mn[:5])

print('[3] 旧例句翻译 worklist')
zh = {}
for p in sorted((ROOT / 'data' / 'translations').glob('sent_zh_*.json')):
    zh.update(json.loads(p.read_text(encoding='utf-8')))
for p in sorted(WORK.glob('tr_out_*.json')):
    zh.update({str(k): v for k, v in
               json.loads(p.read_text(encoding='utf-8')).items()})
wl_total = 0
wl_miss = 0
for line in (ROOT / 'logs' / 'sentence_worklist.tsv').read_text(
        encoding='utf-8').splitlines():
    parts = line.split('\t')
    if len(parts) < 5:
        continue
    wl_total += 1
    if not zh.get(parts[0]):
        wl_miss += 1
print(f'  worklist {wl_total} 句, 缺翻译 {wl_miss}')

print('[4] 游戏 DB (wcpFullEng.db)')
con = sqlite3.connect(str(FULL))
con.execute('PRAGMA busy_timeout=30000')
cur = con.cursor()
tables = [r[0] for r in cur.execute(
    "SELECT name FROM sqlite_master WHERE type='table'")]
print('  tables:', tables)
cols = [r[1] for r in cur.execute('PRAGMA table_info(pron)')]
print('  pron cols:', cols)
pron = dict(cur.execute('SELECT word, ukPhonic FROM pron').fetchall())
print(f'  pron 总词数 {len(pron)}')
miss_pron = [w for w in need if w not in pron]
print(f'  词表词不在 pron: {len(miss_pron)}', miss_pron[:10])
empty_phon = [w for w in need if w in pron and not (pron[w] or '').strip()]
print(f'  pron 里 ukPhonic 为空: {len(empty_phon)}', empty_phon[:10])
s2 = dict(cur.execute(
    'SELECT word, COUNT(*) FROM sentence2 GROUP BY word').fetchall())
s2cnt = Counter(min(s2.get(w, 0), 3) for w in need)
print(f'  sentence2: 词表词有例句 {sum(1 for w in need if w in s2)}, '
      f'总行 {sum(s2.values())}')
print(f'  词表词例句数分布(上限3): {dict(sorted(s2cnt.items()))}')
less3 = [w for w in need if s2.get(w, 0) < 3]
print(f'  DB 中 <3 条例句的词: {len(less3)}')
con.close()

print('[5] 外接词库 (wcpOnlyWord.db)')
con = sqlite3.connect(str(ONLY))
cur = con.cursor()
tables2 = [r[0] for r in cur.execute(
    "SELECT name FROM sqlite_master WHERE type='table'")]
print('  tables:', tables2)
pron2 = dict(cur.execute('SELECT word, ukPhonic FROM pron').fetchall())
print(f'  pron 总词数 {len(pron2)}')
miss2 = [w for w in need if w not in pron2]
print(f'  词表词不在 pron: {len(miss2)}', miss2[:10])
empty2 = [w for w in need if w in pron2 and not (pron2[w] or '').strip()]
print(f'  ukPhonic 为空: {len(empty2)}')
con.close()

print('[6] 音频目录 (packs/ja/audio/word)')
if VOC.exists():
    files = list(VOC.glob('*.mp3')) + list(VOC.glob('*.wav'))
    names = {f.stem for f in files}
    have = sum(1 for w in need if w in names)
    print(f'  音频文件 {len(files)}, 词表命中 {have}/{len(need)}, '
          f'缺 {len(need) - have}')
    miss_a = [w for w in need if w not in names]
    print('  缺音频示例:', miss_a[:8])
else:
    print('  目录不存在:', VOC)

print('[7] 导入 xlsx 读音列抽检')
from openpyxl import load_workbook
for name in ('JLPT_N3.xlsx', 'IT用语.xlsx', '四字熟语.xlsx', '常用汉字.xlsx'):
    p = OUT / 'import' / name
    if not p.exists():
        p = OUT / 'import' / 'setb' / name
    if not p.exists():
        print(f'  {name}: 文件不存在')
        continue
    wb = load_workbook(p, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=1, max_row=200, values_only=True))
    ncols = max(len(r) for r in rows)
    # C 列 (idx2) 读音填充率
    c_fill = sum(1 for r in rows if len(r) > 2 and r[2])
    print(f'  {name}: 行采样 {len(rows)}, 列数 {ncols}, C列(读音)填充 '
          f'{c_fill}/{len(rows)}')
    wb.close()
print('=' * 60)
