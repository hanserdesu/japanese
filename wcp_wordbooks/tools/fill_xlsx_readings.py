# -*- coding: utf-8 -*-
"""把 output/import 下所有 xlsx 的 C 列(读音)补满。

C 列为空时按 gen_words_NN.json 的 reading 填; 纯假名词读音=词本身。
只改 C 列空单元格, 其余列与格式不动 (文件均由 openpyxl 生成, 往返安全)。
"""
import json
import re
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
WORK = ROOT / 'work'
IMPORT = ROOT / 'output' / 'import'
PURE_KANA = re.compile(r'^[\u3040-\u30ffー・]+$')

reading = {}
for n in range(40):
    p = WORK / f'gen_words_{n:02d}.json'
    if p.exists():
        for w, v in json.loads(p.read_text(encoding='utf-8')).items():
            rd = (v.get('reading') or '').strip()
            if not rd and PURE_KANA.match(w):
                rd = w
            reading[w] = rd

total_fill = 0
for p in sorted(IMPORT.rglob('*.xlsx')):
    wb = load_workbook(p)
    ws = wb.active
    fills = 0
    blanks = 0
    for row in ws.iter_rows(min_row=1):
        a = row[0].value
        c = row[2].value if len(row) > 2 else None
        if c is None or str(c).strip() == '':
            rd = reading.get(a)
            if rd:
                row[2].value = rd
                fills += 1
            else:
                blanks += 1
    if fills:
        wb.save(p)
    wb.close()
    total_fill += fills
    print(f'{p.relative_to(IMPORT)}: 补 {fills}, 剩空 {blanks}')
print('共补', total_fill)
